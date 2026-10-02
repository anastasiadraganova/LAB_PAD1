"""Отчёт в Word: experiments/REPORT.md (+ графики) → report/Отчёт_RAG_Солнечная_система.docx

    python experiments/build_docx.py

Поддерживается подмножество Markdown, которое используется в REPORT.md:
заголовки #..###, абзацы, **жирный**, `код`, списки "-" и "1.", таблицы |…|,
картинки ![подпись](figures/x.png), блоки ``` ```.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parent
REPORT = ROOT / "REPORT.md"
OUT = ROOT.parent / "report" / "Отчёт_RAG_Солнечная_система.docx"
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_inline(p, text: str) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            r = p.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(10)
        elif part.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            p.add_run(m.group(1))
        else:
            p.add_run(part)


def title_page(doc) -> None:
    for text, size, bold in [
        ("Министерство науки и высшего образования Российской Федерации", 12, False),
        ("[Название университета]", 12, False), ("", 12, False), ("", 12, False),
        ("Отчёт по лабораторной работе", 16, True),
        ("«Разработка собственной RAG-системы»", 16, True),
        ("Тема: Солнечная система (русская Википедия)", 14, False), ("", 12, False), ("", 12, False),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.font.size, r.bold = Pt(size), bold
    for text in ("Выполнил(а): студент(ка) группы [группа]", "[ФИО]", "", "Проверил: [ФИО преподавателя]"):
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p = doc.add_paragraph("\n\n\n[Город], 2026")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()


def add_table(doc, lines: list[str]) -> None:
    rows = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in lines
            if not re.match(r"^\|?\s*:?-{2,}", ln.strip())]
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row[:len(rows[0])]):
            p = t.cell(i, j).paragraphs[0]
            add_inline(p, cell)
            for r in p.runs:
                r.font.size = Pt(9)
                r.bold = r.bold or i == 0
    doc.add_paragraph()


def main() -> None:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(12)
    for s in doc.sections:
        s.left_margin, s.right_margin = Cm(2.5), Cm(1.5)
    title_page(doc)

    lines = REPORT.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("```"):
            j = i + 1
            while j < len(lines) and not lines[j].startswith("```"):
                j += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(lines[i + 1:j]))
            r.font.name, r.font.size = "Consolas", Pt(8)
            i = j + 1
            continue
        if ln.startswith("|"):
            j = i
            while j < len(lines) and lines[j].startswith("|"):
                j += 1
            add_table(doc, lines[i:j])
            i = j
            continue
        m = re.match(r"^(#{1,3})\s+(.*)", ln)
        img = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)", ln)
        if m:
            doc.add_heading(m.group(2), level=len(m.group(1)) - 1 if len(m.group(1)) > 1 else 0)
        elif img:
            path = (REPORT.parent / img.group(2)).resolve()
            if path.exists():
                doc.add_picture(str(path), width=Cm(15.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph(img.group(1))
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.runs[0].italic = True
        elif re.match(r"^\s*[-*]\s+", ln):
            add_inline(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*]\s+", "", ln))
        elif re.match(r"^\s*\d+\.\s+", ln):
            add_inline(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+\.\s+", "", ln))
        elif ln.startswith(">"):
            p = doc.add_paragraph()
            add_inline(p, ln.lstrip("> "))
            for r in p.runs:
                r.italic = True
        elif ln.strip() and ln.strip() != "---":
            add_inline(doc.add_paragraph(), ln)
        i += 1

    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
